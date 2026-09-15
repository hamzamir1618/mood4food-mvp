"""
Build the v2 dish dataset (Phase 1).

Joins the handoff CSV with the sourcing project's review flags and OCR evidence,
then applies, in order: the owner's worksheet answers, the automated category and
ingredient pass, ingredient-derived allergens and diet flags, USDA-based
nutrition, the automated price check, serving counts and quarantine. Writes
data/dishes_v2.csv and data/build_report.md. No input file is modified.

    python -m pipeline.build_dataset

Quarantined dishes stay in the output with their reason; the seeder keeps them
out of recommendations. The rules are recorded in docs/PHASE1_DATA_DECISIONS.md.
"""

import csv
import json
import re
import sys
from collections import Counter
from datetime import datetime, timezone
from statistics import median

from pipeline import paths, prices, servings
from pipeline.areas import restaurant_locations
from pipeline.classify_categories import PROMPT_VERSION, load_cache
from pipeline.ingredients import (
    SPELLING_TO_NAMES,
    VOCABULARY,
    derive_diet_flags,
    detect_ingredients,
    implied_coating,
)
from pipeline.nutrition import estimate, implausible
from pipeline.sources import dish_uid, load_handoff, load_master, load_ocr_candidates, norm
from pipeline.usda_reference import load_reference
from taste_enrichment import enrich_taste_profiles

TASTE = ("taste_sweet", "taste_salty", "taste_sour", "taste_bitter", "taste_umami", "taste_spice")
NOT_RECOMMENDED = {"beverages", "add_ons"}

# Facts the project owner confirmed that no data source records.
OWNER_CONFIRMATIONS = {
    dish_uid(
        "Terrazza", "Avocado, Bacon And Egg Toast"
    ): "owner confirmed the bacon is turkey (2026-09-14)",
    dish_uid(
        "Terrazza", "Bacon, Sausage & Cheese Omelette"
    ): "owner confirmed the bacon is turkey (2026-09-14)",
    dish_uid("Sakura", "Marinated Beef Yaki"): (
        "owner confirmed the restaurant uses a halal mirin substitute (2026-09-14)"
    ),
    dish_uid("Umai Pan Asian Cuisine", "Beef Tataki"): (
        "owner confirmed the restaurant uses a halal mirin substitute (2026-09-14)"
    ),
}

COLUMNS = [
    "dish_uid",
    "restaurant_name",
    "restaurant_address",
    "restaurant_lat",
    "restaurant_lng",
    "restaurant_area",
    "location_precision",
    "dish_name",
    "raw_dish_name",
    "name_status",
    "category",
    "category_source",
    "category_before",
    "price_rs",
    "price_status",
    "price_note",
    "serves_min",
    "serves_max",
    "serves_source",
    "price_per_person",
    "included_items",
    "ingredients",
    "ingredients_named",
    "ingredients_typical",
    "ingredients_basis",
    "allergens",
    "allergens_known",
    "is_vegan",
    "is_vegetarian",
    "is_halal",
    "halal_note",
    "calories",
    "protein_g",
    "carbs_g",
    "fat_g",
    "nutrition_confidence",
    "nutrition_defaults",
    "nutrition_flag",
    *TASTE,
    "taste_source",
    "review_status",
    "entry_method",
    "source",
    "source_date",
    "quarantined",
    "quarantine_reason",
]

# Names that state the spice level outright. The source's taste values sometimes contradict
# them ("Chicken Pepperoni (Non Spicy)" at 0.8), and the name is the better evidence. A
# karahi is cooked with green chilli, so an unspiced one is a data error, except a Shinwari
# karahi, which is traditionally made with salt and tomato. "Hot" is left out: a hot gulab
# jamun is served hot, not spiced.
MILD_NAME = re.compile(r"\b(non[- ]?spicy|not spicy|no spice|mild)\b", re.I)
SPICED_NAME = re.compile(
    r"\b(spicy|chilli|chili|mirchi|jalapeno|jalapeño|peri[- ]?peri|sriracha|szechuan|"
    r"sichuan|schezwan)\b",
    re.I,
)
KARAHI_NAME = re.compile(r"\bkarahi\b", re.I)
SHINWARI_NAME = re.compile(r"\bshinwari\b", re.I)
MILD_SPICE_CAP = 0.1
SPICED_FLOOR = 0.5


def spice_from_name(name: str, spice) -> float | None:
    """A corrected spice value when the dish's name contradicts the recorded one, else None."""
    value = float(spice or 0)
    if MILD_NAME.search(name):
        return MILD_SPICE_CAP if value > MILD_SPICE_CAP else None
    spiced = SPICED_NAME.search(name) or (
        KARAHI_NAME.search(name) and not SHINWARI_NAME.search(name)
    )
    return SPICED_FLOOR if spiced and value < SPICED_FLOOR else None


PLATTER_SHEETS = ("Platters & combos", "Half & full (optional)")
KEY_COL = "Dish key (don't edit)"
REMOVE = "remove"


# ── Worksheets ───────────────────────────────────────────────────────────────
def _sheet_rows(path, sheet: str) -> list[dict]:
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if sheet not in wb.sheetnames:
        return []
    rows = wb[sheet].iter_rows(values_only=True)
    header = [str(h or "") for h in next(rows, [])]
    return [dict(zip(header, r)) for r in rows]


def _cell_text(value) -> str:
    """A worksheet cell as text. Excel turns '2-3' into a date; read it back as a range."""
    if value is None:
        return ""
    if isinstance(value, datetime):
        return f"{value.month}-{value.day}"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def read_platter_answers() -> tuple[dict, int]:
    answers, total = {}, 0
    if not paths.PLATTER_REVIEW.exists():
        return answers, total
    for sheet in PLATTER_SHEETS:
        for r in _sheet_rows(paths.PLATTER_REVIEW, sheet):
            if not r.get(KEY_COL):
                continue  # formatted but empty rows below the list
            total += 1
            cells = {
                "serves": _cell_text(r.get("SERVES — you fill")),
                "items": _cell_text(r.get("INCLUDED ITEMS — you fill")),
                "notes": _cell_text(r.get("NOTES — you fill")),
            }
            if not any(cells.values()):
                continue
            # "remove" in any answer cell means the owner wants the dish gone. It is
            # quarantined like an OCR discard, and the word is never read as an item list.
            if any(v.lower() == REMOVE for v in cells.values()):
                answers[r[KEY_COL]] = {"remove": True}
            else:
                answers[r[KEY_COL]] = cells
    return answers, total


def read_ocr_decisions() -> tuple[dict, set]:
    decisions, listed = {}, set()
    for path in sorted(paths.REVIEW_DIR.glob(paths.OCR_REVIEW_GLOB)):
        if not path.exists():
            continue
        for r in _sheet_rows(path, "OCR damage"):
            key = r.get(KEY_COL)
            if not key:
                continue
            listed.add(key)
            decision = _cell_text(r.get("DECISION — keep / fix / discard")).lower()
            if decision in ("keep", "fix", "discard"):
                fixed = _cell_text(r.get("FIXED NAME — if fix")) or _cell_text(
                    r.get("Suggested clean name")
                )
                decisions[key] = {"decision": decision, "fixed_name": fixed}
    return decisions, listed


# ── Helpers ──────────────────────────────────────────────────────────────────
def _parse_list(value) -> list[str]:
    if not value:
        return []
    try:
        parsed = json.loads(value)
        return (
            [str(x).lower() for x in parsed] if isinstance(parsed, list) else [str(parsed).lower()]
        )
    except (ValueError, TypeError):
        return [str(value).lower()]


def nutrition_confidence(est: dict | None, basis: str) -> str:
    if est is None:
        return "none"
    if "wheat flour" in est["defaults_used"]:
        return "low"  # no main ingredient identified; the serving is mostly a default
    if basis == "named ingredients" and not est["defaults_used"]:
        return "high"
    return "medium"


def _review_status(master_row: dict) -> str:
    flag = master_row.get("human_confirmed")
    if flag == "True":
        return "human_confirmed"
    if flag == "False":
        return "auto_imported"
    return master_row.get("entry_method") or "unknown"


FILLER = re.compile(
    r"\b(half dozen|dozen|half|full|quarter|single|regular|small|medium|large|"
    r"pcs|pc|pieces|piece|served)\b"
)


def core_name(text: str) -> str:
    """
    A dish or item name without counts, sizes or punctuation, for matching an owner's
    item list against the menu: '4pcs Malai Boti (Half)' -> 'malai boti'.
    """
    words = re.sub(r"[^a-z ]", " ", str(text).lower())
    return " ".join(FILLER.sub(" ", words).split())


def menu_index(rows: list[dict], master: dict) -> dict[tuple[str, str], set[str]]:
    """Each dish's named ingredients, keyed by restaurant and core name."""
    index: dict[tuple[str, str], set[str]] = {}
    for r in rows:
        desc = master.get(r["dish_uid"], {}).get("dish_description") or ""
        key = (r["restaurant_name"], core_name(r["dish_name"]))
        index.setdefault(key, set()).update(detect_ingredients(f"{r['dish_name']} {desc}"))
    return index


def listed_items(restaurant: str, items: str, menu: dict) -> tuple[list[str], bool]:
    """
    The ingredients of an owner's item list, and whether every item on it was recognised.

    An item that is also on the restaurant's own menu brings that dish's ingredients with
    it: "Philadelphia Maki" names no fish, but the menu's Philadelphia Maki does.
    """
    found: set[str] = set()
    all_known = True
    for part in re.split(r"[,\n+&]|\band\b|\bwith\b", items or "", flags=re.I):
        core = core_name(part)
        if not core:
            continue
        ings = set(detect_ingredients(part)) | menu.get((restaurant, core), set())
        found |= ings
        all_known = all_known and bool(ings)
    return [n for n in VOCABULARY if n in found], all_known


# ── Build ────────────────────────────────────────────────────────────────────
def build() -> tuple[list[dict], dict]:
    rows = load_handoff()
    master = load_master()
    candidates = load_ocr_candidates()
    cache = load_cache()
    reference = load_reference()
    platter, platter_total = read_platter_answers()
    ocr_decisions, ocr_listed = read_ocr_decisions()
    medians = prices.restaurant_medians(rows)
    typical_price = servings.typical_single_prices(rows)
    menu = menu_index(rows, master)
    places = restaurant_locations(rows)

    # Taste: the existing seed-time enrichment, run on copies so it cannot alter anything else.
    taste_rows = [dict(r) for r in rows]
    taste_stats = enrich_taste_profiles(taste_rows)

    out = []
    for r, tr in zip(rows, taste_rows):
        uid, key = r["dish_uid"], f"{r['restaurant_name']} | {r['dish_name']}"
        m = master.get(uid, {})
        review_status = _review_status(m)
        raw = (
            (m.get("source_evidence") or "").strip()
            if m.get("entry_method") == "ocr_reviewed"
            else ""
        )
        desc = (m.get("dish_description") or "").strip()
        llm = cache.get(uid) if cache.get(uid, {}).get("prompt_version") == PROMPT_VERSION else None

        # Name: owner decisions first; otherwise flag what either check found
        name, name_status = r["dish_name"], "ok"
        decision = ocr_decisions.get(key)
        if decision and decision["decision"] == "fix" and decision["fixed_name"]:
            name, name_status = decision["fixed_name"], "fixed_by_owner"
        elif decision:
            name_status = {
                "discard": "discarded_by_owner",
                "keep": "kept_by_owner",
                "fix": "flagged_pending_review",
            }[decision["decision"]]
        elif key in ocr_listed or (llm and llm.get("name_damaged")):
            name_status = "flagged_pending_review"

        # Category
        if llm and llm["category"] != "unknown":
            category, category_source = llm["category"], f"llm:{llm['model']}"
        else:
            category, category_source = (
                r["category"],
                ("handoff (llm: unknown)" if llm else "handoff"),
            )

        # Ingredients: named in the dish text or the owner's item list, plus the typical ones
        # from the automated pass
        answer = platter.get(key, {})
        from_items, items_known = listed_items(r["restaurant_name"], answer.get("items", ""), menu)
        in_text = detect_ingredients(" ".join(filter(None, [name, desc])))
        named = [n for n in VOCABULARY if n in set(in_text) | set(from_items)]
        typical = list(llm["ingredients"]) if llm else []
        # Spellings the vocabulary learned after the pass ran ("mozzarella cheese") are
        # recovered from what the pass returned, without asking again.
        for spelling in llm.get("ingredients_rejected", []) if llm else []:
            for n in SPELLING_TO_NAMES.get(str(spelling).strip().lower(), ()):
                if n not in typical:
                    typical.append(n)
        # Restaurants here are halal by default, so pork or alcohol counts only when the dish
        # itself names it. The automated pass sometimes adds them to dumplings or marinades.
        ignored_haram = [i for i in typical if VOCABULARY[i].haram and i not in named]
        typical = [i for i in typical if i not in ignored_haram]
        found = set(named) | set(typical)
        # The pass sometimes files a main dish as an add-on (jumbo prawns, qeema naan).
        # A dish that names meat or seafood is a meal, so it keeps its previous category.
        names_meat = any(VOCABULARY[n].animal in ("meat", "fish", "shellfish") for n in named)
        if category == "add_ons" and names_meat:
            category = r["category"] if r["category"] not in NOT_RECOMMENDED else "other"
            category_source = "handoff (a dish that names meat or seafood is not an add-on)"
        ingredients = [n for n in VOCABULARY if n in found]
        if named:
            basis = (
                "named + typical ingredients" if set(typical) - set(named) else "named ingredients"
            )
        else:
            basis = "typical ingredients" if typical else "none"
        # A name that means a fried coating (nuggets, broast, fish & chips) adds wheat flour for
        # allergens and exclusions only: breading is a small share of a serving, so it stays
        # out of the nutrition estimate.
        coating = implied_coating(name, desc, ingredients)
        for_nutrition = ingredients
        ingredients = [n for n in VOCABULARY if n in found or n in coating]

        # Allergens and diet flags. With no ingredients we know nothing, and say so:
        # allergens unknown, not vegan, not vegetarian (a dish with an unknown recipe is
        # never offered as safe); halal stays the local default.
        legacy_allergens = _parse_list(r["allergens"])
        if ingredients:
            flags = derive_diet_flags(ingredients)
            allergens = sorted(set(flags["allergens"]) | set(legacy_allergens))
            allergens_known, is_halal = True, flags["is_halal"]
            # Vegan and vegetarian are claims about the whole dish, so they need the whole dish
            # assessed: a name that only mentions rice says nothing about the egg in fried rice.
            # An owner's item list counts only when every item on it was recognised.
            whole_dish = bool(typical) or (bool(answer.get("items")) and items_known)
            is_vegan = flags["is_vegan"] and whole_dish
            is_vegetarian = flags["is_vegetarian"] and whole_dish
        else:
            allergens, allergens_known, is_vegan, is_vegetarian, is_halal = (
                sorted(legacy_allergens),
                False,
                False,
                False,
                True,
            )
        if uid in OWNER_CONFIRMATIONS:
            is_halal, halal_note = True, OWNER_CONFIRMATIONS[uid]
        else:
            halal_note = (
                "no pork or alcohol detected" if is_halal else "pork or alcohol named in the dish"
            )

        # Nutrition
        est = estimate(for_nutrition, category, reference)
        confidence = nutrition_confidence(est, basis)

        # Price
        price = float(r["price_rs"] or 0)
        status, note = prices.price_status(
            price, review_status, raw, candidates.get(norm(raw), []) if raw else []
        )
        gross = prices.gross_error(price, medians.get(r["restaurant_name"], 0.0), f"{name} {raw}")

        # Servings
        serv = servings.resolve(
            answer.get("serves"), name, desc, raw, price, typical_price.get(r["restaurant_name"])
        )

        # Taste, with the spice level corrected where the dish's name states it
        taste = {t: tr.get(t, "") for t in TASTE}
        taste_source = tr.get("taste_source", "")
        spice = spice_from_name(name, taste["taste_spice"])
        if spice is not None:
            taste["taste_spice"], taste_source = spice, "name_rule"

        # Quarantine: kept on file, never recommended
        reasons = []
        if gross:
            reasons.append(f"price: {gross}")
        if name_status == "discarded_by_owner":
            reasons.append("owner discarded the name as OCR damage")
        if answer.get("remove"):
            reasons.append("owner removed the dish (platter worksheet)")

        out.append(
            {
                "dish_uid": uid,
                "restaurant_name": r["restaurant_name"],
                "restaurant_address": r["restaurant_address"],
                "restaurant_lat": r["restaurant_lat"],
                "restaurant_lng": r["restaurant_lng"],
                "restaurant_area": places[r["restaurant_name"]]["area"] or "",
                "location_precision": places[r["restaurant_name"]]["precision"],
                "dish_name": name,
                "raw_dish_name": r["dish_name"],
                "name_status": name_status,
                "category": category,
                "category_source": category_source,
                "category_before": r["category"],
                "price_rs": price,
                "price_status": status,
                "price_note": note,
                **serv,
                "price_per_person": servings.price_per_person(
                    price, serv["serves_min"], serv["serves_max"]
                ),
                "included_items": answer.get("items", ""),
                "ingredients": json.dumps(ingredients),
                "ingredients_named": json.dumps(named),
                "ingredients_typical": json.dumps(typical),
                "ingredients_basis": basis,
                "allergens": json.dumps(allergens),
                "allergens_known": allergens_known,
                "is_vegan": is_vegan,
                "is_vegetarian": is_vegetarian,
                "is_halal": is_halal,
                "halal_note": halal_note,
                "calories": est["calories"] if est else "",
                "protein_g": est["protein_g"] if est else "",
                "carbs_g": est["carbs_g"] if est else "",
                "fat_g": est["fat_g"] if est else "",
                "nutrition_confidence": confidence,
                "nutrition_defaults": json.dumps(est["defaults_used"]) if est else "[]",
                "nutrition_flag": implausible(est),
                **taste,
                "taste_source": taste_source,
                "review_status": review_status,
                "entry_method": m.get("entry_method", ""),
                "source": r["source"],
                "source_date": r["source_date"],
                "quarantined": bool(reasons),
                "quarantine_reason": "; ".join(reasons),
                # report-only fields (not written)
                "_legacy": r,
                "_llm_answered": llm is not None,
                "_ignored_haram": ignored_haram,
                "_coating": bool(coating),
                "_spice_rule": spice is not None,
            }
        )

    context = {
        "llm_answered": sum(1 for o in out if o["_llm_answered"]),
        "platter_answers": len(platter),
        "platter_removed": sum(1 for a in platter.values() if a.get("remove")),
        "platter_total": platter_total,
        "ocr_decisions": len(ocr_decisions),
        "ocr_listed": len(ocr_listed),
        "taste_stats": taste_stats,
    }
    return out, context


def write_dataset(out: list[dict]) -> None:
    paths.DATA_DIR.mkdir(parents=True, exist_ok=True)
    with open(paths.DATASET_CSV, "w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=COLUMNS, extrasaction="ignore")
        w.writeheader()
        w.writerows(out)


# ── Report ───────────────────────────────────────────────────────────────────
def _table(counter: Counter, head: tuple[str, str] = ("", "Dishes")) -> list[str]:
    lines = [f"| {head[0]} | {head[1]} |", "|---|---|"]
    lines += [f"| {k} | {v} |" for k, v in counter.most_common()]
    return lines


def _num(v) -> float | None:
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def _first_ingredients(o: dict, n: int) -> str:
    return ", ".join(json.loads(o["ingredients"])[:n]) or "no ingredients found"


def write_report(out: list[dict], ctx: dict) -> None:
    n = len(out)
    places = {o["restaurant_name"]: (o["restaurant_area"], o["location_precision"]) for o in out}
    precision = Counter(p for _, p in places.values())
    unknown = sorted(name for name, (_, p) in places.items() if p == "unknown")
    no_area = sorted(name for name, (a, p) in places.items() if not a and p != "unknown")
    area_counts = Counter(a for a, _ in places.values() if a)
    spice_fixed = [o for o in out if o["_spice_rule"]]
    L = [
        "# Dataset build report",
        "",
        f"Built {datetime.now(timezone.utc).isoformat(timespec='seconds')} from {n} dishes.",
        "",
        "## Inputs",
        "",
        f"- Automated category and ingredient pass (prompt {PROMPT_VERSION}): "
        f"{ctx['llm_answered']} of {n} dishes answered. "
        "Unanswered dishes keep their old category and have no typical ingredients.",
        f"- Platter worksheet: {ctx['platter_answers']} of {ctx['platter_total']} rows answered, "
        f"{ctx['platter_removed']} of them marked remove.",
        f"- OCR worksheet: {ctx['ocr_decisions']} of {ctx['ocr_listed']} rows decided.",
        f"- Taste enrichment: {dict(ctx['taste_stats'])}",
        f"- Spice corrected from the dish's name: {len(spice_fixed)} dishes, e.g. "
        + ", ".join(f"{o['dish_name']} ({o['taste_spice']})" for o in spice_fixed[:8])
        + ".",
        "",
        "## Review status",
        "",
        *_table(Counter(o["review_status"] for o in out)),
        "",
        "## Names",
        "",
        *_table(Counter(o["name_status"] for o in out)),
        "",
    ]

    # Category
    changed = [o for o in out if o["category"] != o["category_before"]]
    judged = [
        o for o in out if o["category_before"] != "other" and o["category_source"].startswith("llm")
    ]
    agree = sum(1 for o in judged if o["category"] == o["category_before"])
    moved_out = [
        o
        for o in changed
        if o["category"] in NOT_RECOMMENDED and o["category_before"] not in NOT_RECOMMENDED
    ]
    other_now = sum(1 for o in out if o["category"] == "other")
    other_before = sum(1 for o in out if o["category_before"] == "other")
    agree_pct = 100 * agree / max(len(judged), 1)
    L += [
        "## Category",
        "",
        f"- Changed: {len(changed)}. Still 'other': {other_now} (was {other_before}).",
        f"- Agreement with the old category where it wasn't 'other': "
        f"{agree} of {len(judged)} ({agree_pct:.1f}%). The old category was "
        "often the restaurant's cuisine rather than the dish's (burgers and wings at a desi "
        "restaurant filed as desi, drinks at a Turkish one as middle eastern), and the largest "
        "groups of changes below are corrections of that kind.",
        f"- Moved into drinks or add-ons, so no longer recommended: {len(moved_out)}.",
        "- Kept as meals although the pass called them add-ons, because they name meat or "
        f"seafood: {sum(1 for o in out if 'not an add-on' in o['category_source'])}.",
        "",
        *_table(Counter(o["category"] for o in out), ("Category", "Dishes")),
        "",
        "Most common changes:",
        "",
        *_table(
            Counter(f"{o['category_before']} → {o['category']}" for o in changed),
            ("Change", "Dishes"),
        )[:16],
        "",
    ]

    # Ingredients, allergens, diet
    legacy_tagged = sum(1 for o in out if _parse_list(o["_legacy"]["allergens"]))
    known = sum(1 for o in out if o["allergens_known"])
    legacy_vegan = [o for o in out if o["_legacy"]["is_vegan"] == "True"]
    vegan_withdrawn = [o for o in legacy_vegan if not o["is_vegan"]]
    tags = Counter(t for o in out for t in json.loads(o["allergens"]))
    ignored = [o for o in out if o["_ignored_haram"]]
    legacy_veg = sum(1 for o in out if o["_legacy"]["is_vegetarian"] == "True")
    not_halal = [o["raw_dish_name"] for o in out if not o["is_halal"]]
    confirmed = sum(1 for o in out if o["dish_uid"] in OWNER_CONFIRMATIONS)
    coated = [o for o in out if o["_coating"]]
    L += [
        "## Ingredients, allergens and diet",
        "",
        *_table(Counter(o["ingredients_basis"] for o in out), ("Ingredients from", "Dishes")),
        "",
        f"- Allergen data known for {known} of {n} dishes ({100 * known / n:.1f}%). "
        f"Before, {legacy_tagged} ({100 * legacy_tagged / n:.1f}%) carried any allergen tag.",
        f"- Allergen tags: {dict(tags.most_common())}",
        f"- Vegan: {sum(o['is_vegan'] for o in out)} (was {len(legacy_vegan)}). "
        f"Vegetarian: {sum(o['is_vegetarian'] for o in out)} (was {legacy_veg}).",
        f"- Old vegan flags withdrawn: {len(vegan_withdrawn)}, e.g. "
        + ", ".join(
            f"{o['raw_dish_name']} ({_first_ingredients(o, 3)})" for o in vegan_withdrawn[:6]
        ),
        f"- Not halal: {len(not_halal)} — {', '.join(not_halal) or 'none'}. "
        f"Owner confirmations applied: {confirmed}.",
        f"- Pork or alcohol suggested by the automated pass for {len(ignored)} dishes "
        "that don't name it, "
        "and ignored, because restaurants here are halal by default"
        + (": " + ", ".join(o["raw_dish_name"] for o in ignored[:10]) if ignored else "")
        + ".",
        f"- Wheat coating implied by the name (nuggets, broast, fish & chips) and added for "
        f"allergens, not nutrition: {len(coated)} dishes"
        + (", e.g. " + ", ".join(o["dish_name"] for o in coated[:10]) if coated else "")
        + ".",
        "- Vegan and vegetarian are only asserted when the whole dish was assessed "
        "(the automated pass answered, or the owner listed what it includes).",
        "",
    ]

    # Nutrition
    pairs = [(o, _num(o["_legacy"]["calories"]), _num(o["calories"])) for o in out]
    both = [(o, a, b) for o, a, b in pairs if a and b]
    flagged = [o for o in out if o["nutrition_flag"]]
    L += [
        "## Nutrition",
        "",
        *_table(Counter(o["nutrition_confidence"] for o in out), ("Confidence", "Dishes")),
        "",
        f"- Median calories per serving, dishes with both estimates: "
        f"{median(a for _, a, _ in both):.0f} before → "
        f"{median(b for _, _, b in both):.0f} now ({len(both)} dishes)."
        if both
        else "- No dishes have both estimates.",
        f"- Flagged as implausible: {len(flagged)}"
        + (
            ": " + "; ".join(f"{o['raw_dish_name']} ({o['nutrition_flag']})" for o in flagged[:8])
            if flagged
            else "."
        ),
        "",
        "Largest changes:",
        "",
        "| Dish | Ingredients | Before kcal | Now kcal | Now protein g |",
        "|---|---|---|---|---|",
        *[
            f"| {o['raw_dish_name']} | {_first_ingredients(o, 5)} | {a:.0f} | {b:.0f} "
            f"| {o['protein_g']} |"
            for o, a, b in sorted(both, key=lambda t: -abs(t[1] - t[2]))[:10]
        ],
        "",
    ]

    # Price, servings, quarantine
    gross = [o for o in out if o["quarantine_reason"].startswith("price")]
    L += [
        "## Price",
        "",
        *_table(Counter(o["price_status"] for o in out), ("Status", "Dishes")),
        "",
        f"Gross errors ({len(gross)}): "
        + (", ".join(f"{o['raw_dish_name']} (Rs {o['price_rs']:.0f})" for o in gross) or "none")
        + ".",
        "",
        "## Locations",
        "",
        f"- {len(places)} restaurants. Coordinates of their own: {precision['place']}. "
        f"The centre of their sector, so distances are approximate: {precision['area']}. "
        f"Unknown, so never used for a distance: {precision['unknown']}"
        + (f" ({', '.join(unknown)})" if unknown else "")
        + ".",
        "- Areas: " + ", ".join(f"{a} {k}" for a, k in area_counts.most_common()) + ".",
        *([f"- No area in the address: {', '.join(no_area)}."] if no_area else []),
        "",
        "## Servings",
        "",
        *_table(Counter(o["serves_source"] for o in out), ("Source", "Dishes")),
        "",
        "## Quarantine",
        "",
        f"{sum(o['quarantined'] for o in out)} dishes are kept on file but never recommended.",
        "",
        *_table(
            Counter(r for o in out if o["quarantined"] for r in o["quarantine_reason"].split("; ")),
            ("Reason", "Dishes"),
        ),
        "",
    ]
    paths.BUILD_REPORT.write_text("\n".join(L) + "\n", encoding="utf-8")


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    out, ctx = build()
    write_dataset(out)
    write_report(out, ctx)
    print(f"wrote {len(out)} dishes -> {paths.DATASET_CSV}")
    print(f"report -> {paths.BUILD_REPORT}")


if __name__ == "__main__":
    main()
