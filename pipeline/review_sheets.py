"""
Review worksheets for the project owner (Phase 1).

    python -m pipeline.review_sheets --ocr-next

Writes the next data_review/ocr_damage_review_batchN.xlsx: dishes the automated
pass flagged as having a damaged name that no existing OCR worksheet lists yet.
Run it again after more dishes are classified and it writes the next batch with
only the new ones. Every batch has the same columns, and the dataset build reads
them all.

The worksheets are filled in by hand, so this module never overwrites one.
"""

import argparse
import os
import sys
from pathlib import Path

from pipeline import paths
from pipeline.classify_categories import PROMPT_VERSION, load_cache
from pipeline.sources import load_handoff, load_master, load_ocr_candidates, norm

OCR_COLS = [  # (key, header, width, filled in by the owner)
    ("review_id", "ID", 8, False),
    ("confidence", "Confidence", 10, False),
    ("restaurant", "Restaurant", 22, False),
    ("dish", "Current name", 38, False),
    ("suggested", "Suggested clean name", 30, False),
    ("decision", "DECISION — keep / fix / discard", 16, True),
    ("fixed_name", "FIXED NAME — if fix", 30, True),
    ("signal", "Why it was flagged", 22, False),
    ("price_rs", "Price (Rs)", 10, False),
    ("category", "Category", 16, False),
    ("reviewed", "Human-reviewed?", 16, False),
    ("raw_ocr_line", "Raw OCR line", 34, False),
    ("where_to_look", "Where to look", 14, False),
    ("dish_key", "Dish key (don't edit)", 28, False),
]
REVIEWED = {"True": "yes", "False": "no (auto-imported)"}


def _help(batch: int) -> list[str]:
    return [
        f"OCR-damaged dish names — batch {batch}",
        "",
        "These are names the automated check flagged that no earlier worksheet listed.",
        "For each row, set DECISION to one of:",
        "  keep    — the name is fine as it is (the automated check is often over-cautious)",
        "  fix     — the dish is real; type the corrected name in FIXED NAME",
        "  discard — not a real dish; it is excluded from recommendations and kept on file,",
        "            never deleted",
        f"You can also just reply with IDs, e.g. 'discard B{batch}-003, B{batch}-010; "
        "keep the rest'.",
    ]


def _key(r: dict) -> str:
    return f"{r['restaurant_name']} | {r['dish_name']}"


def _write(
    path, title: str, cols, rows: list[dict], help_lines: list[str], decision_col: str
) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = title
    ws.append([c[1] for c in cols])
    for r in rows:
        ws.append([r.get(c[0], "") for c in cols])
    head_fill, you_fill = (
        PatternFill("solid", fgColor="2A201A"),
        PatternFill("solid", fgColor="FFF4C2"),
    )
    for ci, (key, _, width, yours) in enumerate(cols, 1):
        head = ws.cell(row=1, column=ci)
        head.font = Font(bold=True, color="FFFFFF")
        head.fill = head_fill
        head.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[head.column_letter].width = width
        for ri in range(2, len(rows) + 2):
            c = ws.cell(row=ri, column=ci)
            c.alignment = Alignment(wrap_text=key in ("dish", "raw_ocr_line"), vertical="top")
            if yours:
                c.fill = you_fill
                c.number_format = "@"  # keep what is typed as text
            if key == "where_to_look":
                c.value = rows[ri - 2]["look_kind"]
                c.hyperlink = rows[ri - 2]["where_to_look"]
                c.font = Font(color="0563C1", underline="single")
    ws.row_dimensions[1].height = 42
    ws.freeze_panes = "D2"
    if rows:
        ws.auto_filter.ref = ws.dimensions
        letter = ws.cell(row=1, column=[c[0] for c in cols].index(decision_col) + 1).column_letter
        dv = DataValidation(type="list", formula1='"keep,fix,discard"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{letter}2:{letter}{len(rows) + 1}")
    hs = wb.create_sheet("How to fill this in")
    for line in help_lines:
        hs.append([line])
    hs["A1"].font = Font(bold=True, size=13)
    hs.column_dimensions["A"].width = 110
    wb.save(path)


def ocr_workbooks() -> list[Path]:
    """The first OCR worksheet and every later batch, in order."""
    return sorted(paths.REVIEW_DIR.glob(paths.OCR_REVIEW_GLOB))


def _already_listed() -> set[str]:
    from pipeline.build_dataset import KEY_COL, _sheet_rows

    return {r.get(KEY_COL) for path in ocr_workbooks() for r in _sheet_rows(path, "OCR damage")}


def _next_batch() -> tuple[int, Path]:
    n = 2
    while (paths.REVIEW_DIR / f"ocr_damage_review_batch{n}.xlsx").exists():
        n += 1
    return n, paths.REVIEW_DIR / f"ocr_damage_review_batch{n}.xlsx"


def ocr_next_batch() -> tuple[Path | None, int]:
    """Write the next batch of flagged names; returns (path, rows), or (None, 0) if none are new."""
    cache = load_cache()
    master = load_master()
    candidates = load_ocr_candidates()
    listed = _already_listed()
    rows = []
    for r in load_handoff():
        answer = cache.get(r["dish_uid"], {})
        if (
            answer.get("prompt_version") != PROMPT_VERSION
            or not answer.get("name_damaged")
            or _key(r) in listed
        ):
            continue
        m = master.get(r["dish_uid"], {})
        raw = (
            (m.get("source_evidence") or "").strip()
            if m.get("entry_method") == "ocr_reviewed"
            else ""
        )
        found = candidates.get(norm(raw), [])
        look, kind = r["source"], "source link"
        if found and found[0].get("_folder"):
            photo = os.path.join(paths.MENU_PHOTOS, found[0]["_folder"], found[0]["_image"])
            if os.path.exists(photo):
                look, kind = photo, "menu photo"
        rows.append(
            {
                "confidence": "possible",
                "signal": "automated check: name looks damaged",
                "restaurant": r["restaurant_name"],
                "dish": r["dish_name"],
                "suggested": "",
                "price_rs": float(r["price_rs"] or 0),
                "category": answer.get("category", r["category"]),
                "reviewed": REVIEWED.get(m.get("human_confirmed"), "scraped"),
                "raw_ocr_line": raw[:200],
                "where_to_look": look,
                "look_kind": kind,
                "dish_key": _key(r),
            }
        )
    if not rows:
        return None, 0
    rows.sort(key=lambda x: (x["restaurant"].lower(), x["dish"].lower()))
    batch, path = _next_batch()
    for i, x in enumerate(rows, 1):
        x["review_id"] = f"B{batch}-{i:03d}"
    paths.REVIEW_DIR.mkdir(parents=True, exist_ok=True)
    _write(path, "OCR damage", OCR_COLS, rows, _help(batch), "decision")
    return path, len(rows)


def main():
    parser = argparse.ArgumentParser(
        description="Write owner review worksheets (never overwrites existing ones)."
    )
    parser.add_argument(
        "--ocr-next",
        action="store_true",
        help="the next batch of names flagged by the automated pass",
    )
    args = parser.parse_args()
    sys.stdout.reconfigure(encoding="utf-8")
    if args.ocr_next:
        path, n = ocr_next_batch()
        print(f"wrote {n} rows -> {path}" if path else "no newly flagged names; nothing written")
    else:
        parser.print_help()


if __name__ == "__main__":
    main()
